"""Turn uploaded files into clean text chunks the search engine can use.

Supported formats: PDF, Word (.docx), Excel (.xlsx), plain text (.txt),
comma-separated values (.csv), and Markdown (.md).

Large documents are broken into overlapping chunks of a few hundred words so
that a single answer points to a focused, readable passage instead of an
entire manual.
"""

from __future__ import annotations

import csv
import io
import os
import re


# Roughly how big each chunk should be, measured in words, with a little
# overlap so sentences spanning a boundary are not lost.
CHUNK_WORDS = 180
CHUNK_OVERLAP_WORDS = 30


class UnsupportedFileType(Exception):
    pass


def extract_text(filename: str, data: bytes) -> str:
    """Pull the raw text out of an uploaded file based on its extension."""
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        return _from_pdf(data)
    if ext == ".docx":
        return _from_docx(data)
    if ext == ".xlsx":
        return _from_xlsx(data)
    if ext in (".txt", ".md"):
        return data.decode("utf-8", errors="replace")
    if ext == ".csv":
        return _from_csv(data)

    raise UnsupportedFileType(
        f"'{ext}' files are not supported. Please use PDF, DOCX, XLSX, "
        f"TXT, MD, or CSV."
    )


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            # A malformed page should not sink the whole document.
            continue
    return "\n\n".join(pages)


def _from_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    # Include any tables, which often hold spec sheets / part numbers.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _from_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _from_csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    return "\n".join(" | ".join(row) for row in reader if any(row))


def _normalize(text: str) -> str:
    # Collapse runs of whitespace but keep paragraph breaks meaningful.
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def chunk_text(text: str) -> list[str]:
    """Split cleaned text into overlapping, word-bounded chunks."""
    text = _normalize(text)
    if not text:
        return []

    words = text.split()
    if len(words) <= CHUNK_WORDS:
        return [text]

    chunks = []
    step = CHUNK_WORDS - CHUNK_OVERLAP_WORDS
    for start in range(0, len(words), step):
        window = words[start : start + CHUNK_WORDS]
        if not window:
            break
        chunks.append(" ".join(window))
        if start + CHUNK_WORDS >= len(words):
            break
    return chunks


def file_to_chunks(filename: str, data: bytes) -> list[str]:
    """Convenience: extract text from a file and chunk it in one step."""
    return chunk_text(extract_text(filename, data))
